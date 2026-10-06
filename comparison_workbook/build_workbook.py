"""Build MPN_PMF_scRNAseq_Analysis.xlsx in the layout of 'Rampal MPN Analysis.xlsx'."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment

E = "openpyxl"; OUT = "MPN_PMF_scRNAseq_Analysis.xlsx"
COMPS = [("PrePMF_vs_Reactive", "Prefibrotic PMF vs Reactive BM"),
         ("MF2_vs_Reactive",    "Overt PMF, MF-2 vs Reactive BM"),
         ("MF3_vs_Reactive",    "Overt PMF, MF-3 vs Reactive BM"),
         ("MF3_vs_PrePMF",      "MF-3 vs Prefibrotic PMF"),
         ("MF3_vs_Early",       "MF-3 vs Prefibrotic + MF-2"),
         ("MPNvsReactive",      "All PMF (pre, MF2, MF3) vs Reactive BM")]
F  = Font(name="Arial", size=10); FB = Font(name="Arial", size=10, bold=True)
FIN = Font(name="Arial", size=10, color="0000FF")                 # editable inputs
HDR = PatternFill("solid", start_color="D9E1F2"); YEL = PatternFill("solid", start_color="FFFF00")

wb = Workbook()
# Make Arial 10 the workbook default font (font index 0 is what unstyled cells use)
from openpyxl.utils.indexed_list import IndexedList
wb._fonts = IndexedList([Font(name="Arial", size=10)])
ws = wb.active; ws.title = "Summary"

def table_sheet(name, df, numfmt):
    sh = wb.create_sheet(name)
    sh.append(list(df.columns))
    for c in sh[1]: c.font = FB; c.fill = HDR
    for row in df.itertuples(index=False):
        sh.append([None if (isinstance(v, float) and v != v) else v for v in row])
    for col_idx, col in enumerate(df.columns, 1):
        fmt = numfmt.get(col)
        if fmt:
            for (cell,) in sh.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                cell.number_format = fmt
    sh.freeze_panes = "B2"; sh.auto_filter.ref = sh.dimensions
    sh.column_dimensions["A"].width = 18
    return sh

DE_FMT = {"logFC": "0.000", "AveExpr": "0.000", "t": "0.000", "P.Value": "0.00E+00",
          "adj.P.Val": "0.00E+00", "B": "0.000"}
nrows = {}
for key, _ in COMPS:
    df = pd.read_csv(f"de_{key}.csv"); nrows[key] = len(df) + 1
    table_sheet(key, df, DE_FMT)
mat = pd.read_csv("log2cpm_matrix.csv")
table_sheet("Log2CPM", mat, {c: "0.00" for c in mat.columns[1:]})
meta = pd.read_csv("metadata.csv")
table_sheet("metadata", meta, {})
print("data sheets written")

# ---- Compare_Rampal: their MF vs Normal next to ours --------------------------
ram = pd.read_excel("../Rampal MPN Analysis.xlsx", sheet_name="MF_NORM", engine=E)
ram = ram.sort_values("P.Value").drop_duplicates("symbol")          # best probe per gene
a = pd.read_csv("de_MF3_vs_Reactive.csv"); b = pd.read_csv("de_MPNvsReactive.csv")
cmp = (ram[["symbol","logFC","adj.P.Val"]]
       .rename(columns={"logFC":"Rampal_MF_logFC","adj.P.Val":"Rampal_MF_adj.P.Val"})
       .merge(a[["symbol","logFC","P.Value","adj.P.Val"]].rename(columns={
              "logFC":"Ours_MF3vsReactive_logFC","P.Value":"Ours_MF3vsReactive_P.Value",
              "adj.P.Val":"Ours_MF3vsReactive_adj.P.Val"}), on="symbol")
       .merge(b[["symbol","logFC","P.Value"]].rename(columns={
              "logFC":"Ours_MPNvsReactive_logFC","P.Value":"Ours_MPNvsReactive_P.Value"}), on="symbol")
       .sort_values("Rampal_MF_adj.P.Val"))
sh = table_sheet("Compare_Rampal", cmp, {"Rampal_MF_logFC":"0.000","Rampal_MF_adj.P.Val":"0.00E+00",
     "Ours_MF3vsReactive_logFC":"0.000","Ours_MF3vsReactive_P.Value":"0.00E+00",
     "Ours_MF3vsReactive_adj.P.Val":"0.00E+00","Ours_MPNvsReactive_logFC":"0.000",
     "Ours_MPNvsReactive_P.Value":"0.00E+00"})
sh["I1"], sh["J1"] = "Rampal_MF_significant", "Direction_vs_ours_MF3"
for c in (sh["I1"], sh["J1"]): c.font = FB; c.fill = HDR
n_cmp = len(cmp) + 1
for r in range(2, n_cmp + 1):
    sh[f"I{r}"] = f"=AND(C{r}<Summary!$C$4,ABS(B{r})>Summary!$C$5)"
    sh[f"J{r}"] = f'=IF(SIGN(B{r})=SIGN(D{r}),"same","opposite")'
sh.auto_filter.ref = f"A1:J{n_cmp}"
for col in "BCDEFGHIJ": sh.column_dimensions[col].width = 16
print("compare rows:", len(cmp))

# ---- Summary (all counts are formulas driven by the blue cutoff cells) ---------
ws["A1"] = "PMF bone marrow scRNA-seq (PRJNA1070224): whole-marrow pseudobulk, limma-voom"; ws["A1"].font = Font(name="Arial", size=12, bold=True)
ws["A3"] = "Cutoffs (edit the blue cells; every count below updates)"; ws["A3"].font = FB
for r, (lab, val, note) in enumerate([
        ("adj.P.Val below", 0.05, "Matches the Rampal workbook Summary (reproduces its ET and MF_PV rows exactly)."),
        ("|logFC| above", 0.585, "log2(1.5) = 1.5-fold; same as the Rampal Summary."),
        ("Nominal P.Value below (exploratory)", 0.01, "Unadjusted; about 1% of genes pass by chance alone.")], start=4):
    ws[f"A{r}"] = lab; ws[f"C{r}"] = val; ws[f"C{r}"].font = FIN; ws[f"C{r}"].fill = YEL
    ws[f"C{r}"].comment = Comment(note, "analysis")
def n(g): return f'COUNTIFS(metadata!$D:$D,"{g}",metadata!$K:$K,TRUE)'
pts = {"PrePMF_vs_Reactive": (n("PrePMF"), n("Reactive")), "MF2_vs_Reactive": (n("MF2"), n("Reactive")),
       "MF3_vs_Reactive": (n("MF3"), n("Reactive")), "MF3_vs_PrePMF": (n("MF3"), n("PrePMF")),
       "MF3_vs_Early": (n("MF3"), f'{n("PrePMF")}+{n("MF2")}'),
       "MPNvsReactive": (f'{n("PrePMF")}+{n("MF2")}+{n("MF3")}', n("Reactive"))}
hdr = ["Comparison","Description","Up","Down","Up (nominal P)","Down (nominal P)","Genes tested","Patients (group vs reference)"]
for j, h in enumerate(hdr, 1):
    c = ws.cell(row=8, column=j, value=h); c.font = FB; c.fill = HDR; c.alignment = Alignment(wrap_text=True)
for i, (key, desc) in enumerate(COMPS, start=9):
    q = f"'{key}'!"
    ws[f"A{i}"], ws[f"B{i}"] = key, desc
    ws[f"C{i}"] = f'=COUNTIFS({q}$G:$G,"<"&$C$4,{q}$C:$C,">"&$C$5)'
    ws[f"D{i}"] = f'=COUNTIFS({q}$G:$G,"<"&$C$4,{q}$C:$C,"<"&-$C$5)'
    ws[f"E{i}"] = f'=COUNTIFS({q}$F:$F,"<"&$C$6,{q}$C:$C,">"&$C$5)'
    ws[f"F{i}"] = f'=COUNTIFS({q}$F:$F,"<"&$C$6,{q}$C:$C,"<"&-$C$5)'
    ws[f"G{i}"] = f"=COUNTA({q}$A:$A)-1"
    g, ref = pts[key]; ws[f"H{i}"] = f'=({g})&" vs "&({ref})'

r0 = 17
ws[f"A{r0}"] = "Agreement with Rampal et al. (MF vs Normal, granulocyte microarray) - sheet Compare_Rampal"; ws[f"A{r0}"].font = FB
rows = [("Genes present in both datasets", "=COUNTA(Compare_Rampal!$A:$A)-1", None),
        ("Rampal MF-significant genes (at the cutoffs above)", "=COUNTIF(Compare_Rampal!$I:$I,TRUE)", None),
        ("  ...same direction in our MF3 vs Reactive", '=COUNTIFS(Compare_Rampal!$I:$I,TRUE,Compare_Rampal!$J:$J,"same")', None),
        ("  ...% same direction (50% expected by chance)", f"=IF(C{r0+2}=0,\"\",C{r0+3}/C{r0+2})", "0.0%"),
        ("Rampal-significant AND our nominal P below cutoff", f'=COUNTIFS(Compare_Rampal!$I:$I,TRUE,Compare_Rampal!$E:$E,"<"&$C$6)', None),
        ("  ...same direction", f'=COUNTIFS(Compare_Rampal!$I:$I,TRUE,Compare_Rampal!$E:$E,"<"&$C$6,Compare_Rampal!$J:$J,"same")', None),
        ("  ...% same direction", f"=IF(C{r0+5}=0,\"\",C{r0+6}/C{r0+5})", "0.0%")]
for k, (lab, f, fmt) in enumerate(rows, start=r0 + 1):
    ws[f"A{k}"] = lab; ws[f"C{k}"] = f
    if fmt: ws[f"C{k}"].number_format = fmt

# ---- Reference blocks (static values; source named in each comment) -----------
r1 = r0 + 9
ws[f"A{r1}"] = "Rampal workbook Summary, for reference"; ws[f"A{r1}"].font = FB
ws[f"A{r1}"].comment = Comment("Source: 'Rampal MPN Analysis.xlsx', sheet Summary (values copied as-is).", "analysis")
for j, h in enumerate(["Comparison", "", "Up", "Down"], 1):
    if h: c = ws.cell(row=r1 + 1, column=j, value=h); c.font = FB; c.fill = HDR
for k, (cmpn, u, d) in enumerate([("ET vs Normal",1834,323),("PV vs Normal",2506,799),
                                   ("MF vs Normal",2473,730),("MF vs PV",30,6)], start=r1 + 2):
    ws[f"A{k}"], ws[f"C{k}"], ws[f"D{k}"] = cmpn, u, d
    for c in (ws[f"C{k}"], ws[f"D{k}"]): c.font = FIN

st = pd.read_csv("../stage14/pseudobulk_summary.csv")
st = st[st.tested == True].sort_values("n_FDR05_autosomal", ascending=False)
r2 = r1 + 8
ws[f"A{r2}"] = "Per-cell-type pseudobulk, MF3 vs Early (edgeR QL, ~ sex + group): DEGs at FDR < 0.05, autosomal"; ws[f"A{r2}"].font = FB
ws[f"A{r2}"].comment = Comment("Source: stage14/pseudobulk_summary.csv in the project repository (static values).", "analysis")
for j, h in enumerate(["Cell type", "Patients MF3 / Early", "DEGs", "Model"], 1):
    c = ws.cell(row=r2 + 1, column=j, value=h); c.font = FB; c.fill = HDR
for k, row in enumerate(st.itertuples(), start=r2 + 2):
    ws[f"A{k}"], ws[f"B{k}"] = row.cell_type, f"{int(row.n_MF3)} / {int(row.n_Early)}"
    ws[f"C{k}"], ws[f"D{k}"] = int(row.n_FDR05_autosomal), row.sex_model
    ws[f"C{k}"].font = FIN

r3 = r2 + 3 + len(st)
ws[f"A{r3}"] = "Notes"; ws[f"A{r3}"].font = FB
for k, t in enumerate([
    "Each patient's QC-passed cells (all cell types) are summed into one whole-marrow profile; limma-voom, design ~ 0 + group + sex.",
    "Sex is adjusted because MF3 is 8 male / 1 female; Reactive marrow (n = 2) is the reference, as Normal is in the Rampal workbook.",
    "SRR27748964 (about 76% late erythroid) is in Log2CPM and metadata but not in the model (metadata column in_model).",
    "Compare_Rampal keeps Rampal's best probe per gene (lowest P.Value). Rampal used granulocyte microarrays; this is whole-marrow scRNA-seq.",
    "Comparison sheets use Rampal's topTable columns (symbol, logFC, AveExpr, t, P.Value, adj.P.Val, B); probeid is replaced by chromosome."],
    start=r3 + 1):
    ws[f"A{k}"] = t
ws.column_dimensions["A"].width = 46; ws.column_dimensions["B"].width = 38
for col in "CDEFG": ws.column_dimensions[col].width = 13
ws.column_dimensions["H"].width = 24
wb.move_sheet("Compare_Rampal", offset=-(len(wb.sheetnames) - 2))
wb.save(OUT); print("saved", OUT, "| sheets:", wb.sheetnames)
